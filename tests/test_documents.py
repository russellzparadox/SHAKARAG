import io
import json

import pytest

from rag.documents import DocIngestError, chunk_sections, extract_any, ingest_document_bytes
from rag.pipeline import RagPipeline


# ---------- extraction ----------

def _xlsx_bytes():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["Item", "Price"])
    ws.append(["Steel", 500])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes():
    from docx import Document as Docx

    d = Docx()
    d.add_heading("Terms", level=1)
    d.add_paragraph("Delivery within 30 days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extract_xlsx():
    kind, sections = extract_any("prices.xlsx", _xlsx_bytes())
    assert kind == "excel"
    assert len(sections) == 1
    assert "## Sheet: Prices" in sections[0]
    assert "Steel | 500" in sections[0]


def test_extract_docx_tables_and_headings():
    kind, sections = extract_any("terms.docx", _docx_bytes())
    assert kind == "word"
    assert "# Terms" in sections[0]
    assert "Delivery within 30 days." in sections[0]


def test_extract_unsupported_and_empty():
    with pytest.raises(DocIngestError):
        extract_any("photo.jpg", b"\xff\xd8\xff")
    with pytest.raises(DocIngestError):
        extract_any("empty.pdf", b"")


def test_chunking_metadata():
    chunks = chunk_sections(["a" * 100], "report.docx", "word", doc_id="abc")
    assert len(chunks) == 1
    c = chunks[0]
    assert c.metadata["kind"] == "document"
    assert c.metadata["source"] == "report.docx"
    assert c.id.startswith("doc:abc:")
    assert "[DOCUMENT WORD]" in c.text


def test_long_text_is_split_with_overlap():
    text = "x" * 4000
    chunks = chunk_sections([text], "big.pdf", "pdf")
    assert len(chunks) >= 3


def test_ingest_into_fake_store(tmp_path):
    from rag.chunking import Chunk

    class FakeStore:
        def __init__(self):
            self.chunks = []

        @property
        def count(self):
            return len(self.chunks)

        def upsert(self, chunks, batch_size=128, progress=None):
            self.chunks.extend(chunks)
            return len(chunks)

    store = FakeStore()
    stats = ingest_document_bytes(store, "prices.xlsx", _xlsx_bytes())
    assert stats["kind"] == "excel"
    assert stats["chunks"] > 0
    assert store.count == stats["chunks"]
    assert all(c.metadata["kind"] == "document" for c in store.chunks)


# ---------- routing ----------

class FakeStore:
    embedder = None

    def __init__(self, hits):
        self._hits = hits

    def query_text(self, text, top_k):
        return self._hits

    def all_table_metas(self):
        return []


def _settings(tmp_path, data_preview=False):
    from rag.config import Settings

    return Settings(
        db_dialect="postgres", db_host="h", db_port=1, db_user="u", db_password="p",
        db_name="d", db_url=None, chroma_dir=tmp_path, collection="c",
        embed_provider="default", openai_api_key=None,
        openai_base_url="http://x", embed_model=None,
        llm_base_url="http://x/v1", llm_model="m", llm_api_key=None,
        llm_temperature=0.0, top_k=4, context_char_budget=4000,
        max_rows=10, statement_timeout_ms=1000,
        sample_values=True, value_sample_max_rows=100000, examples_top_k=2,
        data_preview=data_preview,
    )


def _pipeline(tmp_path, hits, settings=None, llm=None):
    p = RagPipeline.__new__(RagPipeline)
    p.settings = settings or _settings(tmp_path)
    p.store = FakeStore(hits)
    p.llm = llm
    p.dialect = None
    return p


def _hit(kind, source="f.xlsx", table=None):
    meta = {"kind": kind, "source": source}
    if table:
        meta["table"] = table
        meta["schema"] = "dbo"
    return {"metadata": meta, "text": "evidence text", "distance": 0.5}


def test_route_greeting_to_chat(tmp_path):
    p = _pipeline(tmp_path, [])
    assert p._route_intent("hello", [], "") == "chat"


def test_route_data_question_to_sql_even_without_hits(tmp_path):
    p = _pipeline(tmp_path, [])
    assert p._route_intent("how many suppliers are there", [], "") == "sql"


def test_route_doc_only_hits_to_document(tmp_path):
    p = _pipeline(tmp_path, [_hit("document"), _hit("document")])
    assert p._route_intent("what does the contract say", p.store._hits, "") == "document"


def test_mixed_hits_arbitrate_via_llm(tmp_path):
    class RouteLLM:
        def __init__(self, route):
            self.route = route

        def chat_json(self, messages, max_tokens=8192):
            return {"route": self.route}

    hits = [_hit("document"), _hit("schema", table="DimSupplier")]
    p = _pipeline(tmp_path, hits, llm=RouteLLM("document"))
    assert (
        p._route_intent("q", hits, "") == "document"
    )
    p2 = _pipeline(tmp_path, hits, llm=RouteLLM("sql"))
    assert p2._route_intent("q", hits, "") == "sql"


def test_document_answer_cites_sources(tmp_path):
    class AnswerLLM:
        def __init__(self):
            self.last_system = ""

        def chat_json(self, messages, max_tokens=8192):
            return {"route": "document"}

        def chat(self, messages, max_tokens=8192):
            self.last_system = messages[0]["content"]
            return "The delivery term is 30 days [contract.docx]."

    llm = AnswerLLM()
    p = _pipeline(
        tmp_path,
        [_hit("document", source="contract.docx")],
        llm=llm,
    )
    res = p.ask("what is the delivery term?")
    assert res.route == "document"
    assert res.answer.startswith("The delivery term")
    assert res.doc_sources == ["contract.docx"]
    assert "DOCUMENT EXCERPTS" in llm.last_system.upper()


def test_chat_answer_no_sql(tmp_path):
    class ChatLLM:
        def chat(self, messages, max_tokens=8192):
            return "Hi! I can query your database."

    p = _pipeline(tmp_path, [], llm=ChatLLM())
    res = p.ask("hello!")
    assert res.route == "chat"
    assert res.sql is None
    assert "database" in res.answer
